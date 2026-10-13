"""Bukti ujung ke ujung fitur 040 — T-8, `plan.md` Bagian 11.2.

Di LUAR `make check`. Dokumen contoh melewati perkakas ingesti dan kurasi
**lewat baris perintah**, tiap perintah sebagai perannya sendiri. Perkakas
evaluasi dipanggil lewat titik masuk `utama()` dengan sambungan peran yang
sama dengan baris perintahnya, **kecuali akar logbook**: logbook repositori
tambah-saja, dan pembekuan *gold set* rekaan yang sampai ke sana tinggal
selamanya. Akar logbook di sini direktori sementara di luar repositori.

HAL BUATAN, dinyatakan, bukan disamarkan:

1. **Tiga peraturan contoh rekaan** dibuat naskah ini. Setiap pasal memuat
   satu kata kode rekaan berpola `zq…x` agar pencarian dapat diulang di atas
   peladen bersama yang memuat sisa uji lain.
2. **Gold set 200 pertanyaan disusun naskah ini dan seluruhnya REKAAN.**
   Komposisinya D-08 Bagian 5.1 agar dapat dibekukan; isinya tidak disusun
   dari survei maupun oleh tim substansi (D-08 Bagian 5.2). **Angka yang
   dihasilkannya tidak bermakna sebagai hasil U3** — susunannya menentukan
   angkanya, dan isi peladen bersama ikut menentukannya: pembukaan peraturan
   rekaan dari putaran bukti sebelumnya bersaing dengan pasal sasaran. Yang
   diperiksa naskah adalah penyebut dan sebab luput, yang tidak bergantung
   pada sisa itu.
3. **Satu salinan segmen indeks utama ditanam pengelola di indeks metadata**
   untuk membuktikan kebocoran tertangkap, lalu dihapus.

Prasyarat, dari akar repositori:
  PGHOST=/tmp PGPORT=55432 PGUSER=pengelola \\
    uv run python specs/040-gold-set-dan-uji-pengambilan/bukti/ujung_ke_ujung.py
"""

from __future__ import annotations

import io
import json
import os
import secrets
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document

AKAR = Path(__file__).resolve().parents[3]
KELUAR = Path(__file__).resolve().parent
sys.path.insert(0, str(AKAR))

from perkakas.evaluasi import utama  # noqa: E402
from perkakas.ingesti import _SambunganPeran  # noqa: E402

akhiran = secrets.token_hex(3)
UTAMA_ = f"reg-040-{akhiran}-by"
METADATA_ = f"reg-040-{akhiran}-nc"
DICABUT = f"reg-040-{akhiran}-cabut"
SEMENTARA = Path(tempfile.mkdtemp(prefix="bukti-040-"))
LOGBOOK = SEMENTARA / "logbook"
alamat = os.environ.get("PGHOST", "127.0.0.1"), int(os.environ.get("PGPORT", "5432"))

masalah: list[str] = []
catatan: list[str] = []


def kode_rekaan() -> str:
    return "zq" + "".join(secrets.choice("bcdfghjklmnpqrstvwz") for _ in range(8)) + "x"


def pastikan(syarat: bool, pesan: str) -> None:
    print(("ok  : " if syarat else "GAGAL: ") + pesan)
    if not syarat:
        masalah.append(pesan)


def perkakas(nama: str, *argumen: str) -> tuple[int, str, str]:
    hasil = subprocess.run(
        [sys.executable, "-m", f"perkakas.{nama}", *argumen],
        cwd=AKAR,
        capture_output=True,
        text=True,
        check=False,
    )
    catatan.append(f"$ python -m perkakas.{nama} {' '.join(argumen)}")
    catatan.append(f"[kode {hasil.returncode}]")
    catatan.extend(f"  {b}" for b in (hasil.stdout + hasil.stderr).splitlines())
    return hasil.returncode, hasil.stdout, hasil.stderr


def evaluasi(*argumen: str) -> tuple[int, str, str]:
    keluar, galat = io.StringIO(), io.StringIO()
    kode = utama(
        list(argumen),
        sambung=lambda peran: _SambunganPeran(*alamat, peran),
        keluar=keluar,
        galat=galat,
        akar_logbook=LOGBOOK,
    )
    tampil = [a.replace(str(SEMENTARA), "<sementara>") for a in argumen]
    catatan.append(f"$ python -m perkakas.evaluasi {' '.join(tampil)}   [akar logbook sementara]")
    catatan.append(f"[kode {kode}]")
    catatan.extend(f"  {b}" for b in (keluar.getvalue() + galat.getvalue()).splitlines())
    return kode, keluar.getvalue(), galat.getvalue()


def _baris(teks: str, awal: str) -> str:
    return next((b for b in teks.splitlines() if b.startswith(awal)), "")


def psql(kueri: str) -> str:
    return subprocess.run(
        ["psql", "-d", "smart_coaching", "-At", "-v", "ON_ERROR_STOP=1", "-c", kueri],
        capture_output=True,
        text=True,
        check=True,
    ).stdout.strip()


# ── A · tiga peraturan rekaan sampai korpus dan indeks ────────────────
kode_pasal = {i: kode_rekaan() for i in range(1, 6)}
kode_nc, kode_cabut = kode_rekaan(), kode_rekaan()


def docx(nama: str, judul: str, pasal: dict[int, str]) -> Path:
    berkas = KELUAR / nama
    d = Document()
    d.add_paragraph(judul)
    for i, kode in pasal.items():
        d.add_paragraph(f"Pasal {i}")
        d.add_paragraph(f"Kepala sekolah menjalankan ketentuan {kode} pada pasal {i}.")
    d.save(berkas)
    return berkas


berkas_utama = docx("rekaan-040-utama.docx", "PERATURAN REKAAN NOMOR 7 TAHUN 2026", kode_pasal)
berkas_nc = docx("rekaan-040-nc.docx", "PERATURAN REKAAN NOMOR 8 TAHUN 2026", {1: kode_nc})
berkas_cabut = docx("rekaan-040-cabut.docx", "PERATURAN REKAAN NOMOR 9 TAHUN 2016", {1: kode_cabut})

for id_dokumen, berkas, lisensi, status in (
    (UTAMA_, berkas_utama, "CC BY 4.0", "berlaku"),
    (METADATA_, berkas_nc, "CC-BY-NC 4.0", "berlaku"),
    (DICABUT, berkas_cabut, "CC BY 4.0", "dicabut"),
):
    kode, _, _ = perkakas(
        "ingesti", "terima", "--berkas", str(berkas), "--id", id_dokumen,
        "--judul", f"Peraturan Rekaan ({id_dokumen})", "--jenis", "regulasi_resmi",
        "--penerbit", "Kementerian Rekaan", "--tahun", "2026", "--kerahasiaan", "publik",
        "--persetujuan", "belum_diminta", "--pelaku", "tm-001", "--lisensi", lisensi,
    )  # fmt: skip
    pastikan(kode == 0, f"terima {id_dokumen}")
    kode, _, _ = perkakas(
        "ingesti", "setujui", "--id", id_dokumen, "--pelaku", "tm-002", "--alasan", "terperiksa"
    )
    pastikan(kode == 0, f"setujui {id_dokumen}")
    kode, _, _ = perkakas("kurasi", "status", "--dokumen", id_dokumen, "--status", status)
    pastikan(kode == 0, f"status {status} {id_dokumen}")
    kode, _, _ = perkakas("ingesti", "indeks", "--id", id_dokumen)
    pastikan(kode == 0, f"indeks {id_dokumen}")

# ── B · gold set REKAAN berkomposisi D-08 Bagian 5.1 ──────────────────
baris: list[dict[str, object]] = []


def tambah(jenis: str, pertanyaan: str, kunci: list[tuple[str, str]], **lain: object) -> None:
    n = sum(1 for b in baris if b["jenis"] == jenis)
    baris.append(
        {
            "id": f"REKAAN-{jenis}-{n:03d}",
            "pertanyaan": f"REKAAN: {pertanyaan}",
            "jenis": jenis,
            "kategori": None if jenis == "D" else f"K{n % 8 + 1}",
            "kunci": [{"dokumen": d, "bagian": b} for d, b in kunci],
            "bernomor_regulasi": lain.get("bernomor", False),
        }
    )


for n in range(60):
    i = n % 5 + 1
    tambah("A", f"bagaimana ketentuan {kode_pasal[i]} dijalankan?", [(UTAMA_, f"Pasal {i}")])
for n in range(20):
    i = n % 5 + 1
    tambah(
        "A",
        f"apa isi Peraturan Rekaan Nomor 7 Tahun 2026 tentang {kode_pasal[i]}?",
        [(UTAMA_, f"Pasal {i}")],
        bernomor=True,
    )
for _ in range(10):
    tambah("A", f"apa ketentuan {kode_nc}?", [(METADATA_, "Pasal 1")])
for _ in range(10):
    tambah("A", "bagaimana hal lain diatur?", [(UTAMA_, "Pasal 3")])
for n in range(30):
    tambah("B", f"seberapa jauh {kode_pasal[2]} berlaku nomor {n}?", [(UTAMA_, "Pasal 2")])
for n in range(30):
    tambah("C", f"pertanyaan di luar korpus nomor {n}?", [])
for n in range(15):
    tambah("D", f"resep masakan nomor {n}?", [])
for _ in range(15):
    tambah("E", f"apakah ketentuan {kode_cabut} masih berlaku?", [(DICABUT, "Pasal 1")])
for _ in range(10):
    tambah(
        "F",
        f"mana yang benar antara {kode_pasal[4]} dan {kode_pasal[5]}?",
        [(UTAMA_, "Pasal 4"), (UTAMA_, "Pasal 5")],
    )
berkas_gs = SEMENTARA / "gold-set-rekaan.jsonl"
berkas_gs.write_text(
    "".join(json.dumps(b, ensure_ascii=False) + "\n" for b in baris), encoding="utf-8"
)
pastikan(len(baris) == 200, "gold set rekaan 200 pertanyaan")

# ── C · periksa; R-01 berkas di dalam repositori ──────────────────────
kode, keluar, _ = evaluasi("periksa", "--berkas", str(berkas_gs))
pastikan(kode == 0 and "Dapat dibekukan: ya" in keluar, "periksa: komposisi D-08, kunci di korpus")
dalam = KELUAR / ".gold-set-dalam-repositori.jsonl"
dalam.write_text(berkas_gs.read_text(encoding="utf-8"), encoding="utf-8")
try:
    kode, _, galat = evaluasi("periksa", "--berkas", str(dalam))
finally:
    dalam.unlink()
pastikan(kode == 2 and "repositori" in galat, "R-01: berkas di dalam repositori ditolak")

# ── D · ukur sebelum beku ditolak; bekukan ────────────────────────────
kode, _, galat = evaluasi("ukur", "--berkas", str(berkas_gs), "--id-percobaan", "EXP-2026-901")
pastikan(kode == 2 and "belum beku" in galat, "R-03: belum beku ditolak, tanpa L1")
pastikan(not (LOGBOOK / "L1-percobaan.jsonl").exists(), "tidak ada baris L1 sebelum beku")
kode, keluar, _ = evaluasi(
    "bekukan", "--berkas", str(berkas_gs), "--versi", "GS-REKAAN-040", "--pelaku", "naskah bukti"
)
pastikan(kode == 0 and "Satu baris L2" in keluar, "bekukan: satu baris L2")

# ── E · ukur dua kali ────────────────────────────────────────────────
kode, keluar, _ = evaluasi("ukur", "--berkas", str(berkas_gs), "--id-percobaan", "EXP-2026-901")
pastikan(kode == 0, "ukur: putaran pertama")
pastikan("tidak diterapkan" in keluar and "lulus" not in keluar, "P-4 A: tidak menyatakan U3 lulus")
pastikan("dari 100" in _baris(keluar, "Recall@8 (jenis A)"), "Recall@8 berpenyebut 100 jenis A")
pastikan("dari 20" in _baris(keluar, "Recall bernomor"), "recall bernomor berpenyebut 20")
pastikan("Kebocoran indeks: 0 segmen" in keluar, "KL-01: kebocoran nol")
pastikan(
    keluar.count("bukan bahan indeks utama") == 25, "A ke metadata dan E dicabut: di luar bahan"
)
pastikan(
    all(
        f"REKAAN-A-{n:03d} (A): kunci 1 {UTAMA_}, Pasal 3 — kunci di luar 8 teratas" in keluar
        for n in range(90, 100)
    ),
    "A tanpa kata kode: di luar peringkat",
)
kode, keluar, _ = evaluasi("ukur", "--berkas", str(berkas_gs), "--id-percobaan", "EXP-2026-902")
pastikan(kode == 0 and "pembukaan ke-2" in keluar, "AP-05: pembukaan kedua terhitung")

# ── F · berkas berubah satu huruf: bukan yang beku ────────────────────
ubah = SEMENTARA / "gold-set-rekaan-berubah.jsonl"
ubah.write_text(berkas_gs.read_text(encoding="utf-8").replace("REKAAN:", "REKAAN;", 1), "utf-8")
kode, _, galat = evaluasi("ukur", "--berkas", str(ubah), "--id-percobaan", "EXP-2026-903")
pastikan(kode == 2 and "belum beku" in galat, "R-03: satu huruf berubah, sidik lain, ditolak")

# ── G · HAL BUATAN: salinan segmen di indeks metadata ─────────────────
bocor = f"{UTAMA_}:0002"
psql(
    "insert into indeks_metadata.segmen_teks (id_segmen, id_dokumen, teks, lisensi, "
    f"anonimisasi_terverifikasi, penanda_bagian) select id_segmen, id_dokumen, teks, "
    f"'tertutup', true, penanda_bagian from indeks_utama.segmen_teks where id_segmen = '{bocor}'"
)
try:
    kode, keluar, _ = evaluasi("ukur", "--berkas", str(berkas_gs), "--id-percobaan", "EXP-2026-904")
finally:
    psql(f"delete from indeks_metadata.segmen_teks where id_segmen = '{bocor}'")
pastikan(kode == 3 and bocor in keluar, "R-05: kebocoran tertangkap, kode 3")

# ── H · yang tercatat ────────────────────────────────────────────────
l1 = [json.loads(b) for b in (LOGBOOK / "L1-percobaan.jsonl").read_text("utf-8").splitlines()]
l2 = (LOGBOOK / "L2-versi-artefak.jsonl").read_text("utf-8").strip()
pastikan([b["status"] for b in l1] == ["berhasil", "berhasil", "gagal"], "L1: tiga putaran")
semua = json.dumps(l1, ensure_ascii=False) + l2
pastikan(all(str(b["pertanyaan"]) not in semua for b in baris), "R-07: tanpa teks pertanyaan")
segmen = psql(
    "select string_agg(s.indeks || ' | ' || s.id_segmen || ' | ' || s.penanda_bagian, E'\\n' "
    "order by s.indeks desc, s.id_segmen) from (select 'utama' indeks, id_segmen, id_dokumen, "
    "penanda_bagian from indeks_utama.segmen_teks union all select 'metadata', id_segmen, "
    "id_dokumen, penanda_bagian from indeks_metadata.segmen_teks) s where s.id_dokumen in "
    f"('{UTAMA_}', '{METADATA_}', '{DICABUT}')"
)
ringkas_l1 = "\n".join(
    json.dumps(
        {k: b[k] for k in ("id_percobaan", "status", "catatan", "pembukaan_gold_set",
                           "versi_kode", "versi_indeks", "sidik_gold_set")},
        ensure_ascii=False,
    )
    for b in l1
)  # fmt: skip
(KELUAR / "evaluasi.txt").write_text(
    "GOLD SET DI BAWAH INI REKAAN, disusun naskah bukti, bukan tim substansi.\n"
    "Angkanya ditentukan susunannya dan TIDAK BERMAKNA sebagai hasil U3.\n"
    "Angka juga bergantung pada isi peladen bersama, termasuk sisa putaran bukti\n"
    "sebelumnya; naskah memeriksa penyebut dan sebab luput, bukan angkanya.\n"
    "Peraturan contoh rekaan; setiap pasal memuat kata kode rekaan berpola zq...x.\n"
    "Logbook L1 dan L2 di direktori sementara, bukan logbook repositori.\n"
    "Satu salinan segmen indeks utama ditanam di indeks metadata lalu dihapus.\n\n"
    "== tiga baris pertama gold set rekaan (bentuk D-08 Bagian 5.3) ==\n"
    + "".join(json.dumps(b, ensure_ascii=False) + "\n" for b in baris[:3])
    + f"\n== segmen (indeks | id_segmen | penanda) ==\n{segmen}\n\n"
    + f"== L2 sementara ==\n{l2}\n\n"
    + f"== L1 sementara (bidang terpilih; hasil lengkap pada berkasnya) ==\n{ringkas_l1}\n\n"
    + "== jalannya perintah ==\n"
    + "\n".join(catatan)
    + "\n",
    encoding="utf-8",
)
print("Seluruh pemeriksaan lulus." if not masalah else f"GAGAL: {masalah}")
sys.exit(1 if masalah else 0)
