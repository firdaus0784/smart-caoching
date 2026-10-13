"""Bukti sebelum dan sesudah fitur 041 — T-4, `plan.md` Bagian 7.2, R-01.

Di LUAR `make check`. Satu keadaan peladen, dua versi kode:

- **sebelum** — *worktree* git pada commit tepat sebelum T-2 (teks pencarian);
- **sesudah** — *worktree* git pada `HEAD`.

Keduanya menjalankan `python -m perkakas.evaluasi` lewat baris perintah,
dengan peran sungguhan, atas *gold set* beku yang sama. Logbook yang ditulis
adalah logbook *worktree* masing-masing, yang dibuang bersama *worktree*-nya:
logbook repositori tambah-saja dan tidak boleh memuat pembekuan rekaan.

HAL BUATAN, dinyatakan, bukan disamarkan:

1. **Sebelas peraturan rekaan** dibuat naskah ini: sepuluh berjudul serupa —
   nomor 7 sampai 16, tahun dan isi pasal **sama persis** — dan satu yang
   dicabut. Kata kode rekaan berpola `zq…x` pada judul dan isi agar sisa uji
   lain di peladen bersama tidak ikut bersaing.
2. **Gold set 200 pertanyaan disusun naskah ini dan seluruhnya REKAAN.**
   Angkanya ditentukan susunannya dan **tidak bermakna sebagai hasil U3**.
   Yang dibuktikan hanya perbandingan sebelum dan sesudah atas bahan yang sama.

Prasyarat, dari akar repositori:
  PGHOST=/tmp PGPORT=55432 PGUSER=pengelola \\
    uv run python specs/041-identitas-regulasi-pengambilan/bukti/sebelum_sesudah.py
"""

from __future__ import annotations

import json
import secrets
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document

AKAR = Path(__file__).resolve().parents[3]
KELUAR = Path(__file__).resolve().parent
SEMENTARA = Path(tempfile.mkdtemp(prefix="bukti-041-"))
akhiran = secrets.token_hex(3)

masalah: list[str] = []
catatan: list[str] = []


def kode_rekaan() -> str:
    return "zq" + "".join(secrets.choice("bcdfghjklmnpqrstvwz") for _ in range(8)) + "x"


RUN, ISI = kode_rekaan(), kode_rekaan()


def pastikan(syarat: bool, pesan: str) -> None:
    print(("ok  : " if syarat else "GAGAL: ") + pesan)
    if not syarat:
        masalah.append(pesan)


def git(*argumen: str) -> str:
    return subprocess.run(
        ["git", *argumen], cwd=AKAR, capture_output=True, text=True, check=True
    ).stdout.strip()


def perkakas(cwd: Path, label: str, nama: str, *argumen: str) -> tuple[int, str, str]:
    hasil = subprocess.run(
        [sys.executable, "-m", f"perkakas.{nama}", *argumen],
        cwd=cwd,
        capture_output=True,
        text=True,
        check=False,
    )
    tampil = [a.replace(str(SEMENTARA), "<sementara>") for a in argumen]
    catatan.append(f"$ [{label}] python -m perkakas.{nama} {' '.join(tampil)}")
    catatan.append(f"[kode {hasil.returncode}]")
    catatan.extend(f"  {b}" for b in (hasil.stdout + hasil.stderr).splitlines())
    return hasil.returncode, hasil.stdout, hasil.stderr


def judul(n: int, tahun: int = 2026) -> str:
    return f"Peraturan Rekaan {RUN} Nomor {n} Tahun {tahun}"


def id_dokumen(n: int) -> str:
    return f"reg-041-{akhiran}-{n:02d}"


# ── A · worktree sebelum dan sesudah ─────────────────────────────────
t2 = git("log", "--format=%H", "-1", "--grep=^feat(api): teks pencarian memuat judul")
SEBELUM_SHA, SESUDAH_SHA = git("rev-parse", f"{t2}^"), git("rev-parse", "HEAD")
SEBELUM, SESUDAH = SEMENTARA / "sebelum", SEMENTARA / "sesudah"
git("worktree", "add", "--detach", str(SEBELUM), SEBELUM_SHA)
git("worktree", "add", "--detach", str(SESUDAH), SESUDAH_SHA)

try:
    # ── B · sebelas peraturan rekaan lewat baris perintah ─────────────
    def terima(n: int, tahun: int, status: str) -> None:
        berkas = SEMENTARA / f"rekaan-{n:02d}.docx"
        d = Document()
        for p in [
            judul(n, tahun).upper(),
            "Pasal 1",
            f"Kepala sekolah menyusun rencana sarana {ISI}.",
            "Pasal 2",
            f"Komite sekolah menetapkan iuran {ISI}.",
        ]:
            d.add_paragraph(p)
        d.save(berkas)
        kode, _, _ = perkakas(
            AKAR, "repo", "ingesti", "terima", "--berkas", str(berkas), "--id", id_dokumen(n),
            "--judul", judul(n, tahun), "--jenis", "regulasi_resmi", "--penerbit",
            "Kementerian Rekaan", "--tahun", str(tahun), "--kerahasiaan", "publik",
            "--persetujuan", "belum_diminta", "--pelaku", "tm-001", "--lisensi", "CC BY 4.0",
        )  # fmt: skip
        pastikan(kode == 0, f"terima {id_dokumen(n)}")
        for perintah in (
            ("ingesti", "setujui", "--id", id_dokumen(n), "--pelaku", "tm-002", "--alasan", "ok"),
            ("kurasi", "status", "--dokumen", id_dokumen(n), "--status", status),
            ("ingesti", "indeks", "--id", id_dokumen(n)),
        ):
            kode, _, _ = perkakas(AKAR, "repo", *perintah)
            pastikan(kode == 0, f"{perintah[1]} {id_dokumen(n)}")

    for n in range(7, 17):
        terima(n, 2026, "berlaku")
    terima(99, 2016, "dicabut")

    berkas_tolak = SEMENTARA / "rekaan-tanpa-nomor.docx"
    Document().save(berkas_tolak)
    kode, _, galat = perkakas(
        AKAR, "repo", "ingesti", "terima", "--berkas", str(berkas_tolak),
        "--id", f"reg-041-{akhiran}-tolak", "--judul", f"Juknis Rekaan {RUN}",
        "--jenis", "regulasi_resmi", "--penerbit", "Kementerian Rekaan", "--tahun", "2026",
        "--kerahasiaan", "publik", "--persetujuan", "belum_diminta", "--pelaku", "tm-001",
    )  # fmt: skip
    pastikan(kode == 2 and RUN not in galat, "P-3 A: judul regulasi tanpa nomor dan tahun ditolak")

    # ── C · gold set REKAAN ─────────────────────────────────────────
    baris: list[dict[str, object]] = []

    def tambah(jenis: str, pertanyaan: str, kunci: list[tuple[str, str]]) -> None:
        n = sum(1 for b in baris if b["jenis"] == jenis)
        baris.append(
            {
                "id": f"REKAAN-{jenis}-{n:03d}",
                "pertanyaan": f"REKAAN: {pertanyaan}",
                "jenis": jenis,
                "kategori": None if jenis == "D" else f"K{n % 8 + 1}",
                "kunci": [{"dokumen": d, "bagian": b} for d, b in kunci],
                "bernomor_regulasi": jenis in "ABEF",
            }
        )

    for i in range(100):
        n = 7 + i % 10
        if i % 2:
            tambah(
                "A",
                f"apa isi {judul(n)} tentang rencana sarana {ISI}?",
                [(id_dokumen(n), "Pasal 1")],
            )
        else:
            tambah("A", f"bagaimana {judul(n)} mengatur iuran {ISI}?", [(id_dokumen(n), "Pasal 2")])
    for i in range(30):
        n = 7 + i % 10
        tambah("B", f"sejauh mana {judul(n)} mengatur iuran {ISI}?", [(id_dokumen(n), "Pasal 2")])
    for i in range(30):
        tambah("C", f"pertanyaan di luar korpus nomor {i}?", [])
    for i in range(15):
        tambah("D", f"resep masakan nomor {i}?", [])
    for _ in range(15):
        tambah(
            "E",
            f"apakah {judul(99, 2016)} tentang rencana sarana {ISI} berlaku?",
            [(id_dokumen(99), "Pasal 1")],
        )
    for i in range(10):
        n = 7 + i
        tambah(
            "F",
            f"mana yang berlaku dalam {judul(n)}: rencana sarana atau iuran {ISI}?",
            [(id_dokumen(n), "Pasal 1"), (id_dokumen(n), "Pasal 2")],
        )
    gs = SEMENTARA / "gold-set-rekaan.jsonl"
    gs.write_text(
        "".join(json.dumps(b, ensure_ascii=False) + "\n" for b in baris), encoding="utf-8"
    )
    pastikan(len(baris) == 200, "gold set rekaan 200 pertanyaan")

    # ── D · bekukan dan ukur pada kedua worktree ─────────────────────
    laporan: dict[str, str] = {}
    for label, pohon, id_percobaan in (
        ("sebelum", SEBELUM, "EXP-2026-911"),
        ("sesudah", SESUDAH, "EXP-2026-912"),
    ):
        kode, _, _ = perkakas(
            pohon, label, "evaluasi", "bekukan", "--berkas", str(gs),
            "--versi", "GS-REKAAN-041", "--pelaku", "naskah bukti",
        )  # fmt: skip
        pastikan(kode == 0, f"{label}: bekukan pada logbook worktree")
        kode, keluar, _ = perkakas(
            pohon, label, "evaluasi", "ukur", "--berkas", str(gs), "--id-percobaan", id_percobaan
        )
        pastikan(kode == 0, f"{label}: ukur")
        laporan[label] = keluar

    def baris_awal(teks: str, awal: str) -> str:
        return next((b for b in teks.splitlines() if b.startswith(awal)), "")

    rb = {k: baris_awal(v, "Recall bernomor regulasi (jenis A)") for k, v in laporan.items()}
    versi = {k: baris_awal(v, "Versi indeks") for k, v in laporan.items()}

    def persen(b: str) -> float:
        return float(b.split(": ", 1)[1].split("%")[0].replace(",", "."))

    pastikan("pencarian:judul-1" not in versi["sebelum"], "R-04: versi sebelum tanpa penanda")
    pastikan("pencarian:judul-1" in versi["sesudah"], "R-04: versi sesudah berpenanda aturan")
    pastikan(persen(rb["sesudah"]) > persen(rb["sebelum"]), "R-01: recall bernomor naik")
    pastikan(
        rb["sesudah"].startswith("Recall bernomor regulasi (jenis A): 100,0%"), "R-01: seluruhnya"
    )
    pastikan(
        all(
            "dicabut" not in b for b in laporan["sesudah"].splitlines() if b.startswith("Kebocoran")
        ),
        "kebocoran tetap nol",
    )
    pastikan("Kebocoran indeks: 0 segmen" in laporan["sesudah"], "R-05: kebocoran tetap nol")

    (KELUAR / "sebelum-sesudah.txt").write_text(
        "GOLD SET DAN PERATURAN DI BAWAH INI REKAAN, disusun naskah bukti.\n"
        "Angkanya ditentukan susunannya dan TIDAK BERMAKNA sebagai hasil U3; yang\n"
        "dibuktikan hanya perbandingan sebelum dan sesudah atas keadaan peladen yang sama.\n"
        "Sepuluh peraturan berjudul serupa (nomor 7 s.d. 16) berisi pasal yang sama persis.\n"
        "Logbook pembekuan dan putaran adalah logbook worktree, dibuang sesudahnya.\n\n"
        f"Sebelum: {SEBELUM_SHA[:12]} (induk commit T-2)\nSesudah: {SESUDAH_SHA[:12]} (HEAD)\n\n"
        "== ringkasan ==\n"
        + "\n".join(f"{k:8} {versi[k]}\n{'':8} {rb[k]}" for k in ("sebelum", "sesudah"))
        + "\n\n== jalannya perintah ==\n"
        + "\n".join(catatan)
        + "\n",
        encoding="utf-8",
    )
finally:
    for pohon in (SEBELUM, SESUDAH):
        subprocess.run(["git", "worktree", "remove", "--force", str(pohon)], cwd=AKAR, check=False)
    subprocess.run(["git", "worktree", "prune"], cwd=AKAR, check=False)

print("Seluruh pemeriksaan lulus." if not masalah else f"GAGAL: {masalah}")
sys.exit(1 if masalah else 0)
